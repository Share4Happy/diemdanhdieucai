"""
Dịch vụ Quản lý và Xuất Dữ liệu Kiểm Thử Hàng Loạt (Batch Test Service):
- Lưu trữ lịch sử các đợt kiểm thử AI hàng loạt (30 lớp hoặc tùy chọn).
- Tự động sinh file Excel (.xlsx) thống kê chi tiết sĩ số từng lớp.
- Đóng gói trọn gói 1-Click ZIP gồm:
  1. File Excel thống kê sĩ số (thong_ke_si_so.xlsx)
  2. Thư mục chứa toàn bộ ảnh gốc (anh_goc/)
  3. Thư mục chứa toàn bộ ảnh kết quả chấm tròn AI (anh_ket_qua_cham/)
"""

import os
import json
import time
import base64
import zipfile
import shutil
from pathlib import Path
from typing import Optional, List, Dict, Any

from config.settings import settings
from config.logging_config import logger


def _get_base_dir() -> Path:
    base = getattr(settings, "STORAGE_DIR", settings.BASE_DIR / "storage") / "test_batches"
    base.mkdir(parents=True, exist_ok=True)
    return base


def _sanitize_name(name: str) -> str:
    """Loại bỏ ký tự đặc biệt khỏi tên file."""
    for ch in ['\\', '/', ':', '*', '?', '"', '<', '>', '|', ' ']:
        name = name.replace(ch, '_')
    return name


def generate_excel_report(batch_dir: Path, batch_meta: Dict[str, Any], items: List[Dict[str, Any]]) -> Path:
    """Tạo bảng tính Excel (.xlsx) chuyên nghiệp bằng openpyxl."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Thống Kê Sĩ Số"
    ws.views.sheetView[0].showGridLines = True

    # 1. Định nghĩa kiểu dáng
    font_title = Font(name="Calibri", size=15, bold=True, color="FFFFFF")
    font_sub = Font(name="Calibri", size=10, italic=True, color="E2E8F0")
    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_normal = Font(name="Calibri", size=11)
    font_status_ok = Font(name="Calibri", size=11, bold=True, color="166534")
    font_status_absent = Font(name="Calibri", size=11, bold=True, color="991B1B")
    font_status_surplus = Font(name="Calibri", size=11, bold=True, color="9A3412")

    fill_title = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    fill_sub = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_header = PatternFill(start_color="0D9488", end_color="0D9488", fill_type="solid")
    fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    fill_summary = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")

    thin_border_side = Side(border_style="thin", color="CBD5E1")
    border_cell = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
    border_double_bottom = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=Side(border_style="double", color="0F172A"))

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # 2. Tiêu đề chính
    ws.merge_cells("A1:L1")
    cell_a1 = ws["A1"]
    cell_a1.value = "HỆ THỐNG ĐIỂM DANH AI - BÁO CÁO KIỂM THỬ SĨ SỐ HỌC SINH TOÀN TRƯỜNG"
    cell_a1.font = font_title
    cell_a1.fill = fill_title
    cell_a1.alignment = align_center
    ws.row_dimensions[1].height = 36

    # 3. Phụ đề thời gian & thiết bị
    ws.merge_cells("A2:L2")
    cell_a2 = ws["A2"]
    created_time = batch_meta.get("created_at", time.strftime("%d/%m/%Y %H:%M:%S"))
    total_cls = len(items)
    cell_a2.value = f"Đợt kiểm thử: {batch_meta.get('title', 'Đợt Test Hàng Loạt')} | Thời gian: {created_time} | Tổng số lớp: {total_cls} lớp | Mô hình: YOLO26 + Depth Anything V2"
    cell_a2.font = font_sub
    cell_a2.fill = fill_sub
    cell_a2.alignment = align_center
    ws.row_dimensions[2].height = 22

    # 4. Hàng tiêu đề cột (Row 4)
    ws.row_dimensions[3].height = 8  # Dòng trống đệm
    headers = [
        "STT", "Mã Lớp", "Tên Lớp", "Phòng Học",
        "Sĩ Số Chuẩn", "Số HS Đếm Được (AI)", "Vắng / Chênh Lệch",
        "Tỷ Lệ Hiện Diện", "Trạng Thái", "Thời Gian AI (ms)",
        "Tệp Ảnh Gốc", "Tệp Ảnh Chấm Tròn"
    ]
    ws.row_dimensions[4].height = 28

    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = border_cell

    # 5. Dữ liệu các lớp (Row 5 trở đi)
    current_row = 5
    for idx, item in enumerate(items, start=1):
        std_cnt = int(item.get("standard_count") or 40)
        valid_cnt = int(item.get("valid_count") or 0)
        diff = valid_cnt - std_cnt
        rate = round((valid_cnt / max(1, std_cnt)) * 100, 1)
        lat_ms = float(item.get("inference_time_ms") or 0.0)

        if diff == 0:
            status_text = "Đủ Sĩ Số"
            status_font = font_status_ok
        elif diff < 0:
            status_text = f"Vắng {-diff} HS"
            status_font = font_status_absent
        else:
            status_text = f"Thừa {diff} HS"
            status_font = font_status_surplus

        orig_filename = item.get("orig_filename", f"{item.get('code', f'lop_{idx}')}_goc.jpg")
        dot_filename = item.get("dot_filename", f"{item.get('code', f'lop_{idx}')}_cham.jpg")

        row_data = [
            (idx, align_center, font_normal),
            (item.get("code", f"LOP_{idx}"), align_center, font_normal),
            (item.get("name", f"Lớp {idx}"), align_left, font_bold),
            (item.get("room_number", f"P.{idx}"), align_center, font_normal),
            (std_cnt, align_center, font_normal),
            (valid_cnt, align_center, font_bold),
            (f"{'+' if diff > 0 else ''}{diff}", align_center, status_font),
            (f"{rate}%", align_center, font_normal),
            (status_text, align_center, status_font),
            (round(lat_ms, 1), align_right, font_normal),
            (orig_filename, align_left, font_normal),
            (dot_filename, align_left, font_normal),
        ]

        ws.row_dimensions[current_row].height = 20
        is_even = (idx % 2 == 0)

        for col_idx, (val, al, fn) in enumerate(row_data, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = fn
            cell.alignment = al
            cell.border = border_cell
            if is_even:
                cell.fill = fill_zebra

        current_row += 1

    # 6. Hàng Tổng Cộng
    ws.row_dimensions[current_row].height = 24
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    total_label_cell = ws.cell(row=current_row, column=1, value="TỔNG CỘNG TOÀN TRƯỜNG")
    total_label_cell.font = font_bold
    total_label_cell.alignment = align_center

    sum_std = sum(int(item.get("standard_count") or 40) for item in items)
    sum_valid = sum(int(item.get("valid_count") or 0) for item in items)
    sum_diff = sum_valid - sum_std
    avg_rate = round((sum_valid / max(1, sum_std)) * 100, 1)
    avg_lat = round(sum(float(item.get("inference_time_ms") or 0.0) for item in items) / max(1, len(items)), 1)

    summary_values = [
        (5, sum_std, align_center),
        (6, sum_valid, align_center),
        (7, f"{'+' if sum_diff > 0 else ''}{sum_diff}", align_center),
        (8, f"{avg_rate}%", align_center),
        (9, f"Vắng {-sum_diff} HS" if sum_diff < 0 else "Đủ sĩ số", align_center),
        (10, avg_lat, align_right),
        (11, f"{len(items)} ảnh", align_center),
        (12, f"{len(items)} ảnh chấm", align_center),
    ]

    for col_idx in range(1, 13):
        cell = ws.cell(row=current_row, column=col_idx)
        cell.fill = fill_summary
        cell.border = border_double_bottom

    for col_idx, val, al in summary_values:
        cell = ws.cell(row=current_row, column=col_idx, value=val)
        cell.font = font_bold
        cell.alignment = al
        cell.fill = fill_summary
        cell.border = border_double_bottom

    # 7. Tự động căn chỉnh độ rộng cột
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = 0
        for cell in col:
            val_str = str(cell.value or "")
            if cell.row in (1, 2):
                continue
            max_len = max(max_len, len(val_str))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    excel_path = batch_dir / f"thong_ke_si_so_{batch_dir.name}.xlsx"
    wb.save(excel_path)
    return excel_path


def save_batch_test(payload: Dict[str, Any]) -> Dict[str, Any]:
    """
    Lưu trữ toàn bộ đợt kiểm thử hàng loạt:
    - Trích xuất và lưu ảnh gốc vào anh_goc/
    - Trích xuất và lưu ảnh kết quả vào anh_ket_qua_cham/
    - Tạo file Excel thong_ke_si_so_{batch_id}.xlsx
    - Lưu file thông tin info.json
    - Đóng gói file ZIP trọn gói sẵn sàng tải về
    """
    batch_id = payload.get("batch_id") or f"batch_{int(time.time())}"
    title = payload.get("title") or f"Đợt kiểm thử hàng loạt {time.strftime('%d/%m/%Y %H:%M')}"
    created_at = payload.get("created_at") or time.strftime("%d/%m/%Y %H:%M:%S")

    base_dir = _get_base_dir()
    batch_dir = base_dir / batch_id
    batch_dir.mkdir(parents=True, exist_ok=True)

    dir_orig = batch_dir / "anh_goc"
    dir_dots = batch_dir / "anh_ket_qua_cham"
    dir_orig.mkdir(parents=True, exist_ok=True)
    dir_dots.mkdir(parents=True, exist_ok=True)

    raw_items = payload.get("items") or []
    processed_items = []

    for idx, item in enumerate(raw_items, start=1):
        clean_code = _sanitize_name(str(item.get("code") or f"Lop_{idx}"))
        orig_fname = f"{clean_code}_goc.jpg"
        dot_fname = f"{clean_code}_cham.jpg"

        # 1. Lưu ảnh gốc
        orig_b64 = item.get("original_b64") or ""
        if orig_b64 and "base64," in orig_b64:
            try:
                b_data = base64.b64decode(orig_b64.split("base64,")[1])
                (dir_orig / orig_fname).write_bytes(b_data)
            except Exception as e:
                logger.warning(f"Không thể lưu ảnh gốc {orig_fname}: {e}")
        elif item.get("original_path"):
            src = Path(item["original_path"])
            if src.exists():
                shutil.copy2(src, dir_orig / orig_fname)

        # 2. Lưu ảnh kết quả chấm tròn
        ann_b64 = item.get("annotated_b64") or item.get("roi_annotated_b64") or ""
        if ann_b64 and "base64," in ann_b64:
            try:
                b_data = base64.b64decode(ann_b64.split("base64,")[1])
                (dir_dots / dot_fname).write_bytes(b_data)
            except Exception as e:
                logger.warning(f"Không thể lưu ảnh chấm {dot_fname}: {e}")
        elif item.get("annotated_path"):
            src = Path(item["annotated_path"])
            if src.exists():
                shutil.copy2(src, dir_dots / dot_fname)

        cls_name = item.get("classroom_name") or item.get("name") or f"Lớp {clean_code}"
        item_copy = {
            "id": item.get("classroom_id") or item.get("id") or idx,
            "classroom_id": item.get("classroom_id") or item.get("id") or idx,
            "code": clean_code,
            "name": cls_name,
            "classroom_name": cls_name,
            "room_number": item.get("room_number") or "",
            "standard_count": int(item.get("standard_count") or 40),
            "valid_count": int(item.get("valid_count") or 0),
            "inference_time_ms": float(item.get("inference_time_ms") or 0.0),
            "original_image": orig_fname,
            "annotated_image": dot_fname,
            "orig_filename": orig_fname,
            "dot_filename": dot_fname,
            "has_orig_img": (dir_orig / orig_fname).exists(),
            "has_dot_img": (dir_dots / dot_fname).exists()
        }
        processed_items.append(item_copy)

    # 3. Tính toán số liệu thống kê tổng hợp
    total_classes = len(processed_items)
    total_standard = sum(it["standard_count"] for it in processed_items)
    total_valid = sum(it["valid_count"] for it in processed_items)
    discrepancy_classes = sum(1 for it in processed_items if it["valid_count"] != it["standard_count"])
    avg_latency = round(sum(it["inference_time_ms"] for it in processed_items) / max(1, total_classes), 1)

    batch_meta = {
        "batch_id": batch_id,
        "title": title,
        "created_at": created_at,
        "total_classes": total_classes,
        "total_standard": total_standard,
        "total_valid": total_valid,
        "discrepancy_classes": discrepancy_classes,
        "avg_latency_ms": avg_latency,
        "excel_filename": f"thong_ke_si_so_{batch_id}.xlsx",
        "zip_filename": f"Bao_Cao_AI_Test_Hang_Loat_{batch_id}.zip"
    }

    # 4. Sinh file Excel
    try:
        generate_excel_report(batch_dir, batch_meta, processed_items)
    except Exception as e:
        logger.error(f"Lỗi khi sinh file Excel cho batch {batch_id}: {e}", exc_info=True)

    # 5. Lưu info.json
    info_data = {
        **batch_meta,
        "items": processed_items
    }
    (batch_dir / "info.json").write_text(json.dumps(info_data, ensure_ascii=False, indent=2), encoding="utf-8")

    # 6. Đóng gói ZIP trước để sẵn sàng tải tức thì
    create_batch_zip(batch_id)

    return {
        "success": True,
        "batch_id": batch_id,
        "batch_meta": batch_meta,
        "message": f"Đã lưu thành công đợt kiểm thử {batch_id} ({total_classes} lớp học)!"
    }


def create_batch_zip(batch_id: str) -> Path:
    """Đóng gói file ZIP chứa Excel + anh_goc/ + anh_ket_qua_cham/."""
    batch_dir = _get_base_dir() / batch_id
    if not batch_dir.exists():
        raise FileNotFoundError(f"Không tìm thấy đợt test {batch_id}")

    zip_path = batch_dir / f"Bao_Cao_AI_Test_Hang_Loat_{batch_id}.zip"

    excel_file = batch_dir / f"thong_ke_si_so_{batch_id}.xlsx"
    dir_orig = batch_dir / "anh_goc"
    dir_dots = batch_dir / "anh_ket_qua_cham"

    with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
        # 1. Thêm file Excel thống kê
        if excel_file.exists():
            zf.write(excel_file, arcname=f"thong_ke_si_so_{batch_id}.xlsx")
            zf.write(excel_file, arcname="thong_ke_si_so.xlsx")

        # 2. Thêm thư mục ảnh gốc
        if dir_orig.exists():
            for img_file in dir_orig.iterdir():
                if img_file.is_file():
                    zf.write(img_file, arcname=f"anh_goc/{img_file.name}")

        # 3. Thêm thư mục ảnh kết quả chấm
        if dir_dots.exists():
            for img_file in dir_dots.iterdir():
                if img_file.is_file():
                    zf.write(img_file, arcname=f"anh_ket_qua_cham/{img_file.name}")

        # 4. Thêm file JSON dữ liệu thô
        info_json = batch_dir / "info.json"
        if info_json.exists():
            zf.write(info_json, arcname="thong_tin_dot_test.json")

    return zip_path


def list_batch_tests() -> List[Dict[str, Any]]:
    """Lấy danh sách các đợt kiểm thử đã lưu, sắp xếp mới nhất lên đầu."""
    base_dir = _get_base_dir()
    results = []

    for item in base_dir.iterdir():
        if not item.is_dir():
            continue
        info_file = item / "info.json"
        if info_file.exists():
            try:
                data = json.loads(info_file.read_text(encoding="utf-8"))
                results.append({
                    "batch_id": data.get("batch_id", item.name),
                    "title": data.get("title", item.name),
                    "created_at": data.get("created_at", ""),
                    "total_classes": data.get("total_classes", 0),
                    "total_standard": data.get("total_standard", 0),
                    "total_valid": data.get("total_valid", 0),
                    "discrepancy_classes": data.get("discrepancy_classes", 0),
                    "avg_latency_ms": data.get("avg_latency_ms", 0.0),
                    "has_zip": (item / f"Bao_Cao_AI_Test_Hang_Loat_{item.name}.zip").exists()
                })
            except Exception:
                pass

    results.sort(key=lambda x: x.get("batch_id", ""), reverse=True)
    return results


def get_batch_test(batch_id: str) -> Optional[Dict[str, Any]]:
    """Lấy chi tiết một đợt kiểm thử."""
    batch_dir = _get_base_dir() / batch_id
    info_file = batch_dir / "info.json"
    if not info_file.exists():
        return None
    try:
        return json.loads(info_file.read_text(encoding="utf-8"))
    except Exception:
        return None
