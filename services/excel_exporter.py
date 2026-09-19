import os
from datetime import datetime
from pathlib import Path
from typing import Optional, List, Dict
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config.settings import settings
from config.logging_config import logger
from database.db_session import SessionLocal
from database.models import AttendanceSession, AttendanceDetail, Classroom

class ExcelExporter:
    """
    Module tự động hóa xuất báo cáo Excel điểm danh tổng hợp cho cả 30 lớp học
    sử dụng định dạng chuẩn và màu sắc cảnh báo trực quan.
    """

    @staticmethod
    def style_cell(cell, font=None, fill=None, alignment=None, border=None):
        if font: cell.font = font
        if fill: cell.fill = fill
        if alignment: cell.alignment = alignment
        if border: cell.border = border

    def generate_daily_report(self, session_id: int) -> Optional[Path]:
        """Tạo file Excel báo cáo sĩ số cho phiên điểm danh session_id."""
        db = SessionLocal()
        try:
            session = db.query(AttendanceSession).filter(AttendanceSession.id == session_id).first()
            if not session:
                logger.error(f"Không tìm thấy phiên điểm danh với ID {session_id}")
                return None

            details = (
                db.query(AttendanceDetail)
                .join(Classroom)
                .filter(AttendanceDetail.session_id == session_id)
                .order_by(Classroom.id)
                .all()
            )

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Báo Cáo Điểm Danh"
            ws.views.sheetView[0].showGridLines = True

            # Định nghĩa font & style
            font_title = Font(name="Calibri", size=16, bold=True, color="1B365D")
            font_subtitle = Font(name="Calibri", size=11, italic=True, color="555555")
            font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            font_body = Font(name="Calibri", size=11)
            font_body_bold = Font(name="Calibri", size=11, bold=True)

            fill_header = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
            fill_success = PatternFill(start_color="E2F0D9", end_color="E2F0D9", fill_type="solid") # Xanh lá
            fill_warning = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid") # Vàng cam
            fill_danger = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # Đỏ nhạt
            fill_total = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

            thin_border = Border(
                left=Side(style='thin', color='BFBFBF'),
                right=Side(style='thin', color='BFBFBF'),
                top=Side(style='thin', color='BFBFBF'),
                bottom=Side(style='thin', color='BFBFBF')
            )
            double_bottom_border = Border(
                left=Side(style='thin', color='BFBFBF'),
                right=Side(style='thin', color='BFBFBF'),
                top=Side(style='thin', color='BFBFBF'),
                bottom=Side(style='double', color='1B365D')
            )

            # 1. Tiêu đề trường & Báo cáo
            ws.merge_cells("A1:I1")
            ws["A1"] = "TRƯỜNG TRUNG HỌC PHỔ THÔNG ĐIỀU CẢI"
            self.style_cell(ws["A1"], font=Font(name="Calibri", size=12, bold=True, color="333333"),
                            alignment=Alignment(horizontal="center"))

            ws.merge_cells("A2:I2")
            ws["A2"] = "BÁO CÁO TỔNG HỢP ĐIỂM DANH SĨ SỐ HỌC SINH TỰ ĐỘNG BẰNG AI CAMERA"
            self.style_cell(ws["A2"], font=font_title, alignment=Alignment(horizontal="center"))

            # 2. Thông tin phiên quét
            ws.merge_cells("A3:I3")
            ws["A3"] = f"Ngày quét: {session.scan_date} | Thời điểm chụp: {session.scan_time} | Mã phiên: {session.session_code}"
            self.style_cell(ws["A3"], font=font_subtitle, alignment=Alignment(horizontal="center"))

            # 3. Hộp KPI thống kê nhanh
            rate_overall = (session.total_present / session.total_standard * 100) if session.total_standard > 0 else 0
            kpi_rows = [
                ("Tổng số lớp", f"{len(details)} phòng học", "Tổng sĩ số trường", f"{session.total_standard} học sinh"),
                ("Tổng hiện diện", f"{session.total_present} học sinh", "Tổng học sinh vắng", f"{session.total_absent} học sinh"),
                ("Tỷ lệ chuyên cần", f"{rate_overall:.1f}%", "Trạng thái hệ thống", "AI Hoàn tất 100%")
            ]

            start_row = 5
            for idx, (k1, v1, k2, v2) in enumerate(kpi_rows):
                r = start_row + idx
                ws[f"B{r}"] = k1
                ws[f"C{r}"] = v1
                ws[f"F{r}"] = k2
                ws[f"G{r}"] = v2
                self.style_cell(ws[f"B{r}"], font=font_body_bold, alignment=Alignment(horizontal="right"))
                self.style_cell(ws[f"C{r}"], font=font_body, alignment=Alignment(horizontal="left"))
                self.style_cell(ws[f"F{r}"], font=font_body_bold, alignment=Alignment(horizontal="right"))
                self.style_cell(ws[f"G{r}"], font=font_body, alignment=Alignment(horizontal="left"))

            # 4. Bảng chi tiết 30 lớp
            headers = [
                ("STT", 6),
                ("Khối / Lớp", 16),
                ("Phòng học", 14),
                ("Sĩ số chuẩn", 13),
                ("Hiện diện", 13),
                ("Vắng mặt", 13),
                ("Tỷ lệ (%)", 14),
                ("Tình trạng", 18),
                ("Ghi chú đối chứng", 26)
            ]

            table_header_row = 9
            for col_idx, (header_text, width) in enumerate(headers, 1):
                col_letter = get_column_letter(col_idx)
                cell = ws[f"{col_letter}{table_header_row}"]
                cell.value = header_text
                self.style_cell(cell, font=font_header, fill=fill_header,
                                alignment=Alignment(horizontal="center", vertical="center", wrap_text=True),
                                border=thin_border)
                ws.column_dimensions[col_letter].width = width

            ws.row_dimensions[table_header_row].height = 28

            current_row = table_header_row + 1
            for idx, d in enumerate(details, 1):
                cls = d.classroom
                cname = cls.name if cls else f"Lớp {idx}"
                room = cls.room_number if cls else ""
                rate = (d.present_count / d.standard_count * 100) if d.standard_count > 0 else 0

                # Lựa chọn màu sắc tình trạng
                if d.absent_count == 0:
                    status_text = "Đủ sĩ số (100%)"
                    row_fill = fill_success
                elif d.absent_count <= 2:
                    status_text = f"Vắng {d.absent_count} em"
                    row_fill = fill_warning
                else:
                    status_text = f"Vắng nhiều ({d.absent_count} em)"
                    row_fill = fill_danger

                row_values = [
                    (idx, "center", font_body),
                    (cname, "center", font_body_bold),
                    (room, "center", font_body),
                    (d.standard_count, "center", font_body),
                    (d.present_count, "center", font_body_bold),
                    (d.absent_count, "center", font_body_bold),
                    (f"{rate:.1f}%", "center", font_body),
                    (status_text, "center", font_body_bold),
                    (d.notes or "Khớp ảnh AI", "left", font_body)
                ]

                for col_idx, (val, align, f_style) in enumerate(row_values, 1):
                    col_letter = get_column_letter(col_idx)
                    cell = ws[f"{col_letter}{current_row}"]
                    cell.value = val
                    cell_fill = row_fill if col_idx in [6, 7, 8] else None
                    self.style_cell(cell, font=f_style, fill=cell_fill,
                                    alignment=Alignment(horizontal=align, vertical="center"),
                                    border=thin_border)

                ws.row_dimensions[current_row].height = 22
                current_row += 1

            # 5. Hàng Tổng cộng (Total Row)
            total_row = current_row
            ws[f"A{total_row}"] = "TỔNG CỘNG"
            ws.merge_cells(f"A{total_row}:C{total_row}")
            self.style_cell(ws[f"A{total_row}"], font=font_body_bold, fill=fill_total,
                            alignment=Alignment(horizontal="center", vertical="center"),
                            border=double_bottom_border)

            ws[f"D{total_row}"] = f"=SUM(D10:D{total_row - 1})"
            ws[f"E{total_row}"] = f"=SUM(E10:E{total_row - 1})"
            ws[f"F{total_row}"] = f"=SUM(F10:F{total_row - 1})"
            ws[f"G{total_row}"] = f"=(E{total_row}/D{total_row})*100"
            ws[f"H{total_row}"] = "HOÀN TẤT"
            ws[f"I{total_row}"] = "Đối chiếu tự động AI"

            for col_idx in range(1, 10):
                col_letter = get_column_letter(col_idx)
                cell = ws[f"{col_letter}{total_row}"]
                self.style_cell(cell, font=font_body_bold, fill=fill_total,
                                alignment=Alignment(horizontal="center", vertical="center"),
                                border=double_bottom_border)

            ws.row_dimensions[total_row].height = 26

            # 6. Chữ ký xác nhận
            sign_row = total_row + 3
            ws[f"B{sign_row}"] = "CÁN BỘ PHỤ TRÁCH THIẾT BỊ"
            ws[f"G{sign_row}"] = "HIỆU TRƯỞNG DUYỆT"
            self.style_cell(ws[f"B{sign_row}"], font=font_body_bold, alignment=Alignment(horizontal="center"))
            self.style_cell(ws[f"G{sign_row}"], font=font_body_bold, alignment=Alignment(horizontal="center"))

            # Lưu file Excel
            date_clean = session.scan_date.replace("-", "")
            report_dir = settings.REPORTS_DIR / session.scan_date
            report_dir.mkdir(parents=True, exist_ok=True)
            report_filename = f"BaoCaoDiemDanh_{date_clean}_{session.session_code}.xlsx"
            report_path = report_dir / report_filename

            wb.save(str(report_path))
            logger.info(f"Đã xuất thành công file báo cáo Excel: {report_path}")
            return report_path

        except Exception as e:
            logger.error(f"Lỗi khi xuất file Excel báo cáo: {e}")
            return None
        finally:
            db.close()

    def get_attendance_dataframe(self, session_id: int) -> Optional[pd.DataFrame]:
        """Sử dụng pandas để tổng hợp và xử lý dữ liệu báo cáo điểm danh."""
        db = SessionLocal()
        try:
            details = (
                db.query(AttendanceDetail)
                .join(Classroom)
                .filter(AttendanceDetail.session_id == session_id)
                .order_by(Classroom.id)
                .all()
            )
            if not details:
                return None

            data = []
            for idx, d in enumerate(details, 1):
                cls = d.classroom
                cname = cls.name if cls else f"Lớp {idx}"
                room = cls.room_number if cls else ""
                rate = (d.present_count / d.standard_count * 100) if d.standard_count > 0 else 0
                status = "Đủ sĩ số" if d.absent_count == 0 else f"Vắng {d.absent_count} em"

                data.append({
                    "STT": idx,
                    "Khối / Lớp": cname,
                    "Phòng học": room,
                    "Sĩ số chuẩn": d.standard_count,
                    "Hiện diện": d.present_count,
                    "Vắng mặt": d.absent_count,
                    "Tỷ lệ (%)": round(rate, 1),
                    "Tình trạng": status,
                    "Ghi chú": d.notes or "Khớp nhận diện",
                    "Ảnh gốc": d.raw_image_path,
                    "Ảnh AI": d.annotated_image_path
                })

            df = pd.DataFrame(data)
            return df
        finally:
            db.close()

    def list_all_reports(self) -> List[Dict]:
        """Quét toàn bộ các file báo cáo Excel đã xuất trong thư mục storage/reports/."""
        reports = []
        if not settings.REPORTS_DIR.exists():
            return reports

        for f in settings.REPORTS_DIR.glob("**/*.xlsx"):
            if "latest" in f.parts:
                continue
            stat = f.stat()
            size_kb = round(stat.st_size / 1024, 1)
            mtime = datetime.fromtimestamp(stat.st_mtime).strftime("%Y-%m-%d %H:%M:%S")
            rel_path = f.relative_to(settings.REPORTS_DIR)
            reports.append({
                "filename": f.name,
                "relative_path": str(rel_path).replace("\\", "/"),
                "size_kb": size_kb,
                "created_at": mtime,
                "download_url": f"/storage/reports/{str(rel_path).replace('\\', '/')}"
            })

        # Sắp xếp mới nhất lên đầu
        reports.sort(key=lambda x: x["created_at"], reverse=True)
        return reports

excel_exporter = ExcelExporter()

